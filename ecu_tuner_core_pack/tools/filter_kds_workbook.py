#!/usr/bin/env python3
"""Create the intentionally narrow KDS/SDS workbook derivative in this pack.

Usage: python filter_kds_workbook.py INPUT.xlsx OUTPUT.xlsx

The source workbook is not included. This script retains only four reviewed sheets and writes
upstream cached formula results as literal values; it does not execute formulas or validate data.
"""
from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

import openpyxl

SELECTED_SHEETS = [
    "KDS PID",
    "Supported PIDs Z1000",
    "Supported PIDs Z750",
    "Service Codes",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def filter_workbook(source: Path, destination: Path) -> None:
    source_wb = openpyxl.load_workbook(source, data_only=True, read_only=False, keep_links=False)
    missing = [name for name in SELECTED_SHEETS if name not in source_wb.sheetnames]
    if missing:
        raise ValueError(f"source workbook is missing expected sheet(s): {missing}")

    output_wb = openpyxl.Workbook()
    output_wb.remove(output_wb.active)
    for name in SELECTED_SHEETS:
        source_ws = source_wb[name]
        output_ws = output_wb.create_sheet(name)
        for row in source_ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    target = output_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                    target.number_format = cell.number_format
                    if cell.hyperlink:
                        target.hyperlink = cell.hyperlink.target
        for key, dimension in source_ws.column_dimensions.items():
            if dimension.width is not None:
                output_ws.column_dimensions[key].width = dimension.width
            output_ws.column_dimensions[key].hidden = dimension.hidden
        for key, dimension in source_ws.row_dimensions.items():
            if dimension.height is not None:
                output_ws.row_dimensions[key].height = dimension.height
        output_ws.freeze_panes = source_ws.freeze_panes
        output_ws.auto_filter.ref = source_ws.auto_filter.ref

    output_wb.properties.title = "KDS/SDS diagnostic PID reference — filtered"
    output_wb.properties.subject = "Read-only PID and DTC reference; values are not vehicle-verified"
    output_wb.properties.creator = "Curated from HerrRiebmann/KDS2Bluetooth"
    destination.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(destination)

    check = openpyxl.load_workbook(destination, data_only=False, read_only=True, keep_links=False)
    if check.sheetnames != SELECTED_SHEETS:
        raise RuntimeError(f"unexpected output sheets: {check.sheetnames}")
    formula_cells = [
        f"{ws.title}!{cell.coordinate}"
        for ws in check.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if formula_cells:
        raise RuntimeError(f"formula cells were not flattened: {formula_cells[:5]}")
    print(f"source_sha256={sha256(source)}")
    print(f"output_sha256={sha256(destination)}")
    print(f"output_bytes={destination.stat().st_size}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    filter_workbook(args.source, args.destination)


if __name__ == "__main__":
    main()
